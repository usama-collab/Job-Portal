import asyncio
import json
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock

import httpx
import pytest
from fastapi import HTTPException
from google.genai import errors, types

from app.utils import resume_ai


@pytest.fixture
def client(monkeypatch):
    monkeypatch.setattr(resume_ai.settings, "RESUME_AI_ENABLED", True)
    monkeypatch.setattr(resume_ai.settings, "GEMINI_API_KEY", "AIza-test-secret")
    monkeypatch.setattr(resume_ai.settings, "RESUME_AI_MODEL", "gemini-3.1-flash-lite")
    mock = AsyncMock()
    mock.__aenter__.return_value = mock
    monkeypatch.setattr(resume_ai.genai, "Client", Mock(return_value=SimpleNamespace(aio=mock)))
    return mock


def response(payload, reason="STOP"):
    return types.GenerateContentResponse(
        response_id="response_test",
        candidates=[types.Candidate(finish_reason=reason, content=types.Content(parts=[types.Part(text=payload)]))],
    )


@pytest.mark.parametrize("upstream,expected", [(400, 502), (401, 502), (403, 502), (404, 502), (429, 429), (500, 502), (504, 504)])
def test_upstream_errors_log_diagnostics_without_secrets(client, caplog, upstream, expected):
    http_response = httpx.Response(upstream, headers={"x-request-id": "req_test"})
    body = {"error": {"code": upstream, "status": "INVALID_ARGUMENT", "message": "Invalid schema; AIza-test-secret PrivateCandidate"}}
    client.models.generate_content.side_effect = errors.APIError(upstream, body, http_response)
    with pytest.raises(HTTPException) as caught:
        asyncio.run(resume_ai.extract_structured_resume("PrivateCandidate"))
    assert caught.value.status_code == expected
    assert f"status={upstream}" in caplog.text
    assert "request_id=req_test" in caplog.text
    assert "Invalid schema" in caplog.text
    for secret in ("AIza-test-secret", "PrivateCandidate"):
        assert secret not in caplog.text
        assert secret not in caught.value.detail
    client.__aexit__.assert_awaited_once()


@pytest.mark.parametrize("error,expected", [(httpx.ReadTimeout, 504), (httpx.ConnectError, 502)])
def test_transport_errors(client, caplog, error, expected):
    client.models.generate_content.side_effect = error("private transport body")
    with pytest.raises(HTTPException) as caught:
        asyncio.run(resume_ai.extract_structured_resume("PrivateCandidate"))
    assert caught.value.status_code == expected
    assert error.__name__ in caplog.text
    assert "private transport body" not in caplog.text


@pytest.mark.parametrize("reason", ["MAX_TOKENS", "SAFETY", "RECITATION", "OTHER"])
def test_incomplete_or_refused_response(client, caplog, reason):
    client.models.generate_content.return_value = response("PrivateCandidate", reason)
    with pytest.raises(HTTPException) as caught:
        asyncio.run(resume_ai.extract_structured_resume("PrivateCandidate"))
    assert caught.value.status_code == 422
    assert "response_test" in caplog.text
    assert reason in caplog.text
    assert "PrivateCandidate" not in caplog.text


def test_prompt_blocked(client, caplog):
    client.models.generate_content.return_value = types.GenerateContentResponse(
        prompt_feedback=types.GenerateContentResponsePromptFeedback(block_reason="SAFETY"),
    )
    with pytest.raises(HTTPException) as caught:
        asyncio.run(resume_ai.extract_structured_resume("PrivateCandidate"))
    assert caught.value.status_code == 422
    assert "SAFETY" in caplog.text


def test_success_request_configuration_and_evidence(client):
    client.models.generate_content.return_value = response(json.dumps({
        "skills": [{"name": "Python", "source_excerpt": "Python"}, {"name": "Go", "source_excerpt": "Go"}],
    }))
    result = asyncio.run(resume_ai.extract_structured_resume("Python candidate@example.com"))
    assert [item.name for item in result.skills] == ["Python"]
    assert result.warnings == ["1 item was omitted because supporting resume text could not be verified."]
    kwargs = client.models.generate_content.call_args.kwargs
    assert kwargs["model"] == "gemini-3.1-flash-lite"
    assert kwargs["config"].response_json_schema == resume_ai.AIResumeExtraction.model_json_schema()
    assert kwargs["config"].response_mime_type == "application/json"
    assert kwargs["config"].system_instruction == resume_ai.INSTRUCTIONS
    assert "candidate@example.com" not in kwargs["contents"]
    options = resume_ai.genai.Client.call_args.kwargs
    assert options["vertexai"] is False
    assert options["http_options"].timeout == 40_000
    assert options["http_options"].retry_options.attempts == 1


@pytest.mark.parametrize("payload", ["private-output", '{"projects":[{"name":"PrivateCandidate","url":"private-output"}]}', '{"work_experience":[{"company":"Acme","title":"Engineer","start_date":"2024","end_date":"2020"}]}'])
def test_invalid_model_output_never_logs_validation_inputs(client, caplog, payload):
    client.models.generate_content.return_value = response(payload)
    with pytest.raises(HTTPException) as caught:
        asyncio.run(resume_ai.extract_structured_resume("PrivateCandidate"))
    assert caught.value.status_code == 422
    assert "ValidationError" in caplog.text
    assert "private-output" not in caplog.text
    assert "PrivateCandidate" not in caplog.text


@pytest.mark.parametrize("enabled,key", [(False, "AIza-test"), (True, None), (True, "   ")])
def test_configuration_errors(client, monkeypatch, enabled, key):
    monkeypatch.setattr(resume_ai.settings, "RESUME_AI_ENABLED", enabled)
    monkeypatch.setattr(resume_ai.settings, "GEMINI_API_KEY", key)
    with pytest.raises(HTTPException) as caught:
        asyncio.run(resume_ai.extract_structured_resume("PrivateCandidate"))
    assert caught.value.status_code == 503
    client.models.generate_content.assert_not_awaited()


def test_real_sdk_serializes_schema_and_parses_response(monkeypatch):
    original_client = resume_ai.genai.Client
    monkeypatch.setattr(resume_ai.settings, "RESUME_AI_ENABLED", True)
    monkeypatch.setattr(resume_ai.settings, "GEMINI_API_KEY", "AIza-test-secret")
    monkeypatch.setattr(resume_ai.settings, "RESUME_AI_MODEL", "gemini-3.1-flash-lite")
    requests = []

    def handle(request):
        requests.append(request)
        body = json.loads(request.content)
        assert request.url.path.endswith("models/gemini-3.1-flash-lite:generateContent")
        assert body["generationConfig"]["responseMimeType"] == "application/json"
        schema = body["generationConfig"]["responseJsonSchema"]
        assert set(body["generationConfig"]) == {"responseMimeType", "responseJsonSchema", "maxOutputTokens"}
        assert body["generationConfig"]["maxOutputTokens"] == 6000
        assert schema == resume_ai.AIResumeExtraction.model_json_schema()
        assert '"maxItems"' not in json.dumps(schema)
        assert schema["additionalProperties"] is False
        assert "projects" in schema["properties"]
        return httpx.Response(200, json={
            "candidates": [{"finishReason": "STOP", "content": {"parts": [{"text": "{}"}]}}],
        })

    def factory(**kwargs):
        kwargs["http_options"].httpx_async_client = httpx.AsyncClient(transport=httpx.MockTransport(handle))
        return original_client(**kwargs)

    monkeypatch.setattr(resume_ai.genai, "Client", factory)
    assert asyncio.run(resume_ai.extract_structured_resume("Synthetic resume")).skills == []
    assert len(requests) == 1


def test_provider_schema_preserves_structure_and_application_constraints():
    from app.schemas.resume import ResumeExtraction

    original = ResumeExtraction.model_json_schema()
    wire = resume_ai.AIResumeExtraction.model_json_schema()
    assert original["properties"]["skills"]["maxItems"] == 100
    assert original["$defs"]["Project"]["properties"]["technologies"]["maxItems"] == 30
    assert '"maxItems"' not in json.dumps(wire)
    assert wire["properties"]["projects"]["items"] == {"$ref": "#/$defs/Project"}
    project = wire["$defs"]["Project"]
    assert project["additionalProperties"] is False
    assert project["required"] == ["name", "source_excerpt"]
    for definition in wire["$defs"].values():
        assert "source_excerpt" in definition["required"]
        assert definition["properties"]["source_excerpt"]["type"] == "string"
        assert "anyOf" not in definition["properties"]["source_excerpt"]
    assert "source_excerpt" not in original["$defs"]["Project"]["required"]
    assert project["properties"]["url"]["anyOf"] == [{"maxLength": 2083, "minLength": 1, "type": "string"}, {"type": "null"}]
    assert project["properties"]["name"]["minLength"] == 1
    assert project["properties"]["name"]["maxLength"] == 200
    assert "pattern" in project["properties"]["start_date"]["anyOf"][0]
    assert "title" in wire["$defs"]["WorkExperience"]["properties"]
    assert ResumeExtraction.model_json_schema() == original


@pytest.mark.parametrize("payload", [
    {"skills": [{"name": "Python"}] * 101},
    {"work_experience": [{"company": "Example", "title": "Engineer"}] * 31},
    {"education": [{"institution": "Example"}] * 21},
    {"projects": [{"name": "Demo"}] * 31},
    {"warnings": ["Ambiguous"] * 21},
    {"projects": [{"name": "Demo", "technologies": ["Python"] * 31}]},
    {"skills": [{"name": ""}]},
    {"skills": [{"name": "x" * 201}]},
    {"skills": [{"name": "Python", "source_excerpt": "x" * 301}]},
    {"projects": [{"name": "Demo", "description": "x" * 2001}]},
    {"projects": [{"name": "Demo", "url": "ftp://example.com"}]},
    {"projects": [{"name": "Demo", "url": "https://example.com/" + "x" * 2083}]},
    {"projects": [{"name": "Demo", "start_date": "2024-13"}]},
    {"work_experience": [{"company": "Example", "title": "Engineer", "is_current": True, "end_date": "2024"}]},
    {"projects": [{"name": "Demo", "unexpected": "value"}]},
])
def test_provider_output_still_enforces_strict_application_validation(client, payload):
    client.models.generate_content.return_value = response(json.dumps(payload))
    with pytest.raises(HTTPException) as caught:
        asyncio.run(resume_ai.extract_structured_resume("Synthetic resume: Python"))
    assert caught.value.status_code == 422


def test_resume_style_text_through_parser_and_analysis(client):
    from io import BytesIO
    from docx import Document
    from app.utils.resume_parser import extract_resume_text

    document = Document()
    document.add_paragraph("Alex Example — Software Engineer building web applications.")
    document.add_paragraph("Skills: Ｐｙｔｈｏｎ • Fast\u200bAPI • React | Postgre-\nSQL • Next.js • Docker • Tailwind\u00a0CSS")
    document.add_paragraph("Example Labs — Software Engineer, building APIs and web applications.")
    buffer = BytesIO()
    document.save(buffer)
    source = extract_resume_text(buffer.getvalue(), "synthetic.docx")
    skills = ["Python", "FastAPI", "React", "PostgreSQL", "Next.js", "Docker", "Tailwind CSS"]
    client.models.generate_content.return_value = response(json.dumps({
        "skills": [{"name": name, "source_excerpt": name} for name in skills] + [{"name": "Kubernetes", "source_excerpt": "Kubernetes"}],
        "work_experience": [{"company": "Example Labs", "title": "Software Engineer", "source_excerpt": "Example Labs: Software Engineer"}],
    }))
    result = asyncio.run(resume_ai.extract_structured_resume(source))
    assert [item.name for item in result.skills] == skills
    assert len(result.work_experience) == 1
    assert len(result.warnings) == 1


def test_reproduced_response_without_evidence_remains_rejected(client):
    client.models.generate_content.return_value = response(json.dumps({
        "skills": [{"name": name} for name in ["Python", "FastAPI", "React", "PostgreSQL"]],
    }))
    result = asyncio.run(resume_ai.extract_structured_resume("Python, FastAPI, React, PostgreSQL"))
    assert result.skills == []
    assert result.warnings == ["4 items were omitted because supporting resume text could not be verified."]
