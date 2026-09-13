from io import BytesIO

import pytest
from docx import Document
from pydantic import ValidationError
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.schemas.resume import ExtractedSkill, ResumeExtraction, WorkExperience
from app.utils.resume_ai import _redact_contacts, _validated_evidence
from app.utils.resume_parser import ResumeParseError, _repair_pdf_letter_spacing, extract_resume_text


def _rotated_text_pdf(text: str) -> bytes:
    """Build a text-layer PDF whose page content is rotated like the reported CV."""

    buffer = BytesIO()
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)}),
    })
    content = DecodedStreamObject()
    content.set_data(f"BT /F1 12 Tf 0 1 -1 0 300 100 Tm ({text}) Tj ET".encode("ascii"))
    page[NameObject("/Contents")] = writer._add_object(content)
    writer.write(buffer)
    return buffer.getvalue()


def test_docx_resume_text_is_extracted_in_order():
    buffer = BytesIO()
    document = Document()
    document.add_heading("Jane Engineer")
    document.add_paragraph("Python, FastAPI, PostgreSQL")
    document.add_paragraph("Software Engineer at Acme from 2020 to 2024 building production systems.")
    document.save(buffer)

    text = extract_resume_text(buffer.getvalue(), "resume.docx")

    assert text.index("Jane Engineer") < text.index("Python") < text.index("Acme")


def test_scanned_or_empty_resume_is_rejected():
    buffer = BytesIO()
    document = Document()
    document.add_paragraph("Too short")
    document.save(buffer)

    with pytest.raises(ResumeParseError, match="No usable text"):
        extract_resume_text(buffer.getvalue(), "resume.docx")


def test_empty_multipage_pdf_is_rejected_as_scanned():
    buffer = BytesIO()
    writer = PdfWriter()
    for _ in range(10):
        writer.add_blank_page(width=612, height=792)
    writer.write(buffer)

    with pytest.raises(ResumeParseError, match="Scanned PDFs are not supported"):
        extract_resume_text(buffer.getvalue(), "resume.pdf")


def test_pdf_falls_back_to_plain_extraction_for_rotated_text_layer(caplog):
    resume_text = (
        "Experienced Python engineer building reliable API systems and database "
        "applications for five years."
    )

    with caplog.at_level("INFO", logger="app.utils.resume_parser"):
        text = extract_resume_text(_rotated_text_pdf(resume_text), "resume.pdf")

    assert resume_text in text
    assert "pages=1" in caplog.text
    assert "'selected_mode': 'plain'" in caplog.text
    assert resume_text not in caplog.text


def test_items_without_matching_evidence_are_removed():
    extraction = ResumeExtraction(skills=[
        ExtractedSkill(name="Python", source_excerpt="Python"),
        ExtractedSkill(name="Kubernetes", source_excerpt="Kubernetes"),
    ])

    validated = _validated_evidence(extraction, "Experienced with Python and FastAPI")

    assert [skill.name for skill in validated.skills] == ["Python"]
    assert validated.warnings == ["1 item was omitted because supporting resume text could not be verified."]


def test_structured_dates_must_be_consistent():
    with pytest.raises(ValidationError):
        WorkExperience(company="Acme", title="Engineer", start_date="2024", end_date="2023")


def test_pdf_parser_output_keeps_supported_resume_skills():
    source = extract_resume_text(_rotated_text_pdf(
        "Alex Example. Software Engineer building APIs and web applications. "
        "Skills: Python | FastAPI | React | PostgreSQL | Next.js | Docker."
    ), "synthetic.pdf")
    names = ["Python", "FastAPI", "React", "PostgreSQL", "Next.js", "Docker"]
    excerpt = "Skills: Python, FastAPI, React, PostgreSQL, Next.js, Docker."
    extraction = ResumeExtraction(skills=[
        {"name": name, "source_excerpt": excerpt} for name in names
    ])
    assert [item.name for item in _validated_evidence(extraction, source).skills] == names


def test_rotated_letter_spaced_cv_retains_skills_and_redacts_contacts():
    # The reported PDF emits single spaces between every glyph and wider word
    # gaps. Earlier fixtures covered rotation but missed this text-layer shape.
    original = (
        "Alex Example Full Stack Developer. Skills: Python, FastAPI, React, "
        "PostgreSQL, Docker, Next.js. Email: candidate@example.com "
        "Phone: 0307-1234567. Built web applications and database systems."
    )
    source = extract_resume_text(_rotated_text_pdf(" ".join(original)), "letter-spaced.pdf")
    names = ["Python", "FastAPI", "React", "PostgreSQL", "Docker", "Next.js"]
    result = _validated_evidence(ResumeExtraction(skills=[
        {"name": name, "source_excerpt": name} for name in names
    ] + [{"name": "Kubernetes", "source_excerpt": "Kubernetes"}]), source)
    assert [item.name for item in result.skills] == names
    assert len(result.warnings) == 1
    redacted = _redact_contacts(source)
    assert "candidate@example.com" not in redacted
    assert "0307-1234567" not in redacted
    assert "[email removed]" in redacted
    assert "[phone removed]" in redacted


@pytest.mark.parametrize("text", [
    "Python FastAPI React PostgreSQL",
    "Py thon Fast API Postgre SQL",
    "P y t h o n",
    "C C++ C# .NET Next.js " * 20,
    "A B C D E F G H I J " * 10,  # No separate word boundaries to recover.
])
def test_letter_spacing_repair_does_not_join_ordinary_or_ambiguous_text(text):
    assert _repair_pdf_letter_spacing(text) == text


def test_letter_spacing_repair_preserves_word_boundaries_and_normal_lines():
    spaced = " ".join("Software Engineer building Python and React applications. " * 5)
    source = spaced + "\nNormal text: Py thon must stay separate."
    repaired = _repair_pdf_letter_spacing(source)
    assert "Software   Engineer" in repaired
    assert "Normal text: Py thon must stay separate." in repaired
    assert "SoftwareEngineer" not in repaired
